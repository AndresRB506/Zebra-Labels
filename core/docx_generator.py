from __future__ import annotations

import os
from docx import Document
from docx.shared import Inches
from typing import List, Dict

from .label_layout import fill_label_cell, SLOTS_PER_PAGE

def _sanitize_filename(name: str) -> str:
    invalid = '<>:"/\\|?*'
    for ch in invalid:
        name = name.replace(ch, "_")
    return name.strip()

def _new_labels_table(doc: Document):
    """
    Creates a 2x3 table used as one 'page' of 6 label slots.
    We do NOT insert page breaks; Word will paginate naturally.
    """
    table = doc.add_table(rows=2, cols=3)
    table.autofit = False

    # Approx cell sizing (tweak if you want)
    for row in table.rows:
        for cell in row.cells:
            cell.width = Inches(2.6)

    return table

def _iter_cells_in_slot_order(table):
    # slot order: row0 col0..2, row1 col0..2
    for r in range(2):
        for c in range(3):
            yield table.cell(r, c)

def generate_docx(lot_number: str, color: str, work_orders: List[Dict], sheets: List[Dict], output_dir: str) -> str:
    """
    Generates DOCX:
    - Cover: up to 4 WO labels with QTY = total_qty
    - Sheets: for each sheet, add label per WO with allocation > 0
      (quantity hidden on sheet labels, per your rules)
    - No forced page breaks between tables.
    """
    doc = Document()

    # Ensure output folder exists
    os.makedirs(output_dir, exist_ok=True)

    # ---------- COVER ----------
    cover_table = _new_labels_table(doc)
    cover_cells = list(_iter_cells_in_slot_order(cover_table))

    # Put up to 4 WOs on cover (slots 0..3); remaining slots left blank
    for i, wo in enumerate(work_orders[:4]):
        fill_label_cell(
            cover_cells[i],
            lot_number=lot_number,
            wo=wo,
            qty_override=int(wo.get("total_qty", 1)),
            hide_qty=False
        )

    # ---------- SHEETS ----------
    # We'll fill labels sequentially across 6-slot tables.
    current_table = _new_labels_table(doc)
    cell_iter = iter(_iter_cells_in_slot_order(current_table))

    def next_cell():
        nonlocal current_table, cell_iter
        try:
            return next(cell_iter)
        except StopIteration:
            # Start next table WITHOUT page break
            current_table = _new_labels_table(doc)
            cell_iter = iter(_iter_cells_in_slot_order(current_table))
            return next(cell_iter)

    for sheet in sheets:
        sheet_no = int(sheet["sheet_number"])
        allocations = sheet["allocations"]  # list[(wo_index, qty)]
        for (wo_index, qty) in allocations:
            if qty <= 0:
                continue

            cell = next_cell()
            fill_label_cell(
                cell,
                lot_number=lot_number,
                wo=work_orders[int(wo_index)],
                sheet_number=sheet_no,
                hide_qty=True
            )

    # Save
    filename = _sanitize_filename(f"LOT {lot_number} {color}.docx")
    path = os.path.join(output_dir, filename)
    doc.save(path)
    return path
